"""Seed L-Acoustics PA bundles from the K2/KARA II/L2 package workbook.

Reads /app/backend/seed_data/lacoustics_k2_karaii_l2_packages.xlsx — one sheet
per package. Each sheet has the same layout:

    Row 5  : header row (S.NO. | MAKE | MODEL | DESCRIPTION | QTY | UNIT | UNIT PRICE (INR) | AMOUNT (INR))
    Row 6+ : section headers (MAIN SPEAKER / SUBWOOFERS / …) — col A blank
             component rows                                     — col A = serial no.
    Last   : Total / GST / Grand Total

Components with AMOUNT == "OPTIONAL" (e.g. KS28-COV) are listed as components
but flagged optional; their price is NOT included in `fixed_price_inr` (the
sheet author already excluded them from the Total row).

`fixed_price_inr` is taken from the row labelled "Total:" in column G — this is
the Hi-Tech selling price before 18 % GST, mirroring the DiGiCo workbook.
"""
import asyncio
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import load_workbook

ROOT = Path(__file__).parent
XLSX = ROOT / "seed_data" / "lacoustics_k2_karaii_l2_packages.xlsx"


def _parse_sheet(ws) -> dict:
    """Extract one package from a single worksheet."""
    components = []
    total_inr = None
    for r in range(1, ws.max_row + 1):
        sno = ws.cell(r, 1).value
        col_f = ws.cell(r, 6).value  # "Total:" label sits in col F
        col_g = ws.cell(r, 7).value  # numeric total / unit price

        # Capture the Total: row
        if isinstance(col_f, str) and col_f.strip().rstrip(":").lower() == "total":
            try:
                total_inr = float(col_g)
            except Exception:
                pass
            break

        # Skip header & section-header rows
        if sno is None or not str(sno).strip().isdigit():
            continue

        model = (ws.cell(r, 3).value or "").strip() if ws.cell(r, 3).value else ""
        description = (ws.cell(r, 4).value or "").strip() if ws.cell(r, 4).value else ""
        try:
            qty = int(ws.cell(r, 5).value)
        except Exception:
            continue
        unit_price = ws.cell(r, 7).value  # noqa: F841 (kept for future use)
        amount = ws.cell(r, 8).value
        is_optional = isinstance(amount, str) and amount.strip().upper() == "OPTIONAL"
        _ = unit_price  # explicit no-op

        components.append({
            "model": model or None,
            "description": description or model or "",
            "qty": qty,
            "optional": is_optional,
        })

    return {"components": components, "fixed_price_inr": total_inr}


async def main():
    load_dotenv(ROOT / ".env")
    db = AsyncIOMotorClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]

    if not XLSX.exists():
        raise SystemExit(f"Pricelist not found: {XLSX}")

    wb = load_workbook(XLSX, data_only=True, read_only=False)

    # Idempotent: drop any previous imports from this workbook
    deleted = await db.packages.delete_many({"brand": "L-Acoustics", "source": "k2_karaii_l2_2026"})
    print(f"Cleared {deleted.deleted_count} previously-imported L-Acoustics packages")

    inserted = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parsed = _parse_sheet(ws)
        if not parsed["fixed_price_inr"] or not parsed["components"]:
            print(f"  skipped (no total or no components): {sheet_name}")
            continue
        doc = {
            "id": str(uuid.uuid4()),
            "brand": "L-Acoustics",
            "name": sheet_name.strip(),
            "sku": None,
            "description": f"L-Acoustics turnkey PA bundle ({len(parsed['components'])} components, GST extra).",
            "items": [],
            "components": parsed["components"],
            "fixed_price_inr": round(parsed["fixed_price_inr"], 2),
            "source": "k2_karaii_l2_2026",
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.packages.insert_one(doc)
        inserted += 1
        opt_count = sum(1 for c in parsed["components"] if c.get("optional"))
        print(f"  · {sheet_name:<24} ₹{parsed['fixed_price_inr']:>15,.0f}  ({len(parsed['components'])} comp, {opt_count} optional)")
    print(f"\nInserted {inserted} L-Acoustics packages")


if __name__ == "__main__":
    asyncio.run(main())
