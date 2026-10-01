"""Seed DiGiCo bundle packages from the 2026 pricelist Excel.

Reads /app/backend/seed_data/digico_pricelist_2026.xlsx → sheet 'Packages Working'.

Layout per package (column letters):
  A  S.NO.       — non-empty marks the start of a new package
  B  Description — package name (only on header row)
  C  Model       — SKU of component line (also surface SKU on header row)
  D  Qty         — component qty
  E  GBP unit
  F  GBP amount
  J  CP (42%)    — final INR sell-price for the package (only on header row)

Each header row is followed by 1..N component rows; package ends at the next
header row or at a blank totals row.

Bundle pricing: `fixed_price_inr` is taken from column J — this is the Hi-Tech
selling price (INR) including any bundle discount, so it overrides the
component-sum in the quotation builder.
"""
import asyncio
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import load_workbook

ROOT = Path(__file__).parent
XLSX = ROOT / "seed_data" / "digico_pricelist_2026.xlsx"


def _read_packages():
    wb = load_workbook(XLSX, data_only=True, read_only=False)
    ws = wb["Packages Working"]
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append([ws.cell(r, c).value for c in range(1, 12)])

    pkgs = []
    cur = None
    for r in rows[2:]:  # skip 2 header rows
        sno = r[0]
        # Header row: column A has a serial number
        if sno is not None and str(sno).strip().isdigit():
            if cur:
                pkgs.append(cur)
            name = str(r[1] or "").strip().replace("\n", " ")
            sku = str(r[2] or "").strip() if r[2] else None
            try:
                inr = float(r[9]) if r[9] not in (None, "") else None  # col J
            except Exception:
                inr = None
            cur = {
                "name": name,
                "sku": sku,
                "fixed_price_inr": inr,
                "components": [],
            }
            # The header row's col C is the main surface SKU + col D qty
            if sku and r[3]:
                try:
                    cur["components"].append({
                        "model": sku,
                        "description": sku,
                        "qty": int(r[3]),
                    })
                except Exception:
                    pass
            continue
        # Component row: col C has model, col D has qty
        model = (str(r[2]).strip() if r[2] else None)
        qty = r[3]
        if cur and model and qty:
            try:
                cur["components"].append({
                    "model": model.split(" - ")[0].strip(),
                    "description": model,
                    "qty": int(qty),
                })
            except Exception:
                pass
    if cur:
        pkgs.append(cur)
    return [p for p in pkgs if p["name"] and p["fixed_price_inr"]]


async def main():
    load_dotenv(ROOT / ".env")
    db = AsyncIOMotorClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]

    if not XLSX.exists():
        raise SystemExit(f"Pricelist not found: {XLSX}")

    parsed = _read_packages()
    print(f"Parsed {len(parsed)} packages from {XLSX.name}")

    # Idempotent: drop existing imported DiGiCo packages
    deleted = await db.packages.delete_many({"brand": "DiGiCo", "source": "pricelist_2026"})
    print(f"Cleared {deleted.deleted_count} previously-imported DiGiCo packages")

    inserted = 0
    for pkg in parsed:
        doc = {
            "id": str(uuid.uuid4()),
            "brand": "DiGiCo",
            "name": pkg["name"],
            "sku": pkg["sku"],
            "description": f"Hi-Tech bundle ({len(pkg['components'])} components). Includes any bundle discount.",
            "items": [],  # no product-id linkage for imported bundles
            "components": pkg["components"],
            "fixed_price_inr": round(pkg["fixed_price_inr"], 2),
            "source": "pricelist_2026",
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.packages.insert_one(doc)
        inserted += 1
        print(f"  · {pkg['name'][:60]:<60} ₹{pkg['fixed_price_inr']:>14,.0f}  ({len(pkg['components'])} comp.)")
    print(f"\nInserted {inserted} DiGiCo packages")


if __name__ == "__main__":
    asyncio.run(main())
